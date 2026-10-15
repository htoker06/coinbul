"""Raporları Excel (.xlsx) ve PDF olarak üretir."""
import io
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import db
from raporlar import hucre_metni

PARA_BICIMI = '#,##0.00 "TL"'


def excel(rapor, okul_adi=""):
    wb = Workbook()
    ws = wb.active
    ws.title = rapor["baslik"][:31]
    kalin = Font(bold=True)
    ince = Side(style="thin", color="BBBBBB")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)

    satir = 1
    if okul_adi:
        ws.cell(satir, 1, okul_adi).font = Font(bold=True, size=13)
        satir += 1
    ws.cell(satir, 1, rapor["baslik"]).font = Font(bold=True, size=12)
    satir += 1
    for b in rapor["bilgi"]:
        ws.cell(satir, 1, b)
        satir += 1
    ws.cell(satir, 1, "Rapor tarihi: " + datetime.now().strftime("%d.%m.%Y %H:%M"))
    satir += 2

    baslik_satiri = satir
    for i, (ad, _) in enumerate(rapor["sutunlar"], 1):
        c = ws.cell(satir, i, ad)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = kenar
    satir += 1

    ilk_veri = satir
    for veri in rapor["satirlar"]:
        for i, ((_, tip), deger) in enumerate(zip(rapor["sutunlar"], veri), 1):
            c = ws.cell(satir, i, _excel_deger(deger, tip))
            _bicimle(c, tip)
            c.border = kenar
        satir += 1
    son_veri = satir - 1

    if rapor["toplam"] is not None:
        for i, ((_, tip), deger) in enumerate(zip(rapor["sutunlar"], rapor["toplam"]), 1):
            harf = get_column_letter(i)
            if tip in ("para", "sayi") and deger != "" and son_veri >= ilk_veri:
                c = ws.cell(satir, i, f"=SUM({harf}{ilk_veri}:{harf}{son_veri})")
            else:
                c = ws.cell(satir, i, _excel_deger(deger, tip))
            _bicimle(c, tip)
            c.font = kalin
            c.border = kenar
            c.fill = PatternFill("solid", fgColor="DDEBF7")

    for i, (ad, tip) in enumerate(rapor["sutunlar"], 1):
        uzunluk = max([len(ad)] + [len(hucre_metni(v[i - 1], tip)) for v in rapor["satirlar"]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(uzunluk + 3, 10), 45)
    if rapor["satirlar"]:
        ws.auto_filter.ref = (f"A{baslik_satiri}:{get_column_letter(len(rapor['sutunlar']))}"
                              f"{son_veri}")
    ws.freeze_panes = ws.cell(baslik_satiri + 1, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def _excel_deger(deger, tip):
    if deger == "" or deger is None:
        return ""
    if tip == "para" and isinstance(deger, int):
        return deger / 100
    if tip == "tarih":
        return db.tarih_yaz(deger)
    if tip == "yuzde" and isinstance(deger, int):
        return deger / 100
    return deger


def _bicimle(c, tip):
    if tip == "para":
        c.number_format = PARA_BICIMI
        c.alignment = Alignment(horizontal="right")
    elif tip == "yuzde":
        c.number_format = "0%"
        c.alignment = Alignment(horizontal="right")
    elif tip in ("sayi", "tarih"):
        c.alignment = Alignment(horizontal="center")


_FONT_HAZIR = False


def _font_kaydet():
    """Türkçe karakter (ğ, ş, ı, İ) için TTF font kaydeder."""
    global _FONT_HAZIR
    if _FONT_HAZIR:
        return
    burasi = os.path.dirname(os.path.abspath(__file__))
    adaylar = [
        (os.path.join(burasi, "fonts", "DejaVuSans.ttf"),
         os.path.join(burasi, "fonts", "DejaVuSans-Bold.ttf")),
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for normal, kalin in adaylar:
        if os.path.exists(normal) and os.path.exists(kalin):
            pdfmetrics.registerFont(TTFont("Okul", normal))
            pdfmetrics.registerFont(TTFont("Okul-Kalin", kalin))
            pdfmetrics.registerFontFamily("Okul", normal="Okul", bold="Okul-Kalin")
            _FONT_HAZIR = True
            return
    raise RuntimeError("PDF için Türkçe destekli font bulunamadı (fonts/DejaVuSans.ttf).")


def pdf(rapor, okul_adi="", hazirlayan="", onaylayan="Okul Müdürü"):
    _font_kaydet()
    tampon = io.BytesIO()
    genis = len(rapor["sutunlar"]) > 6
    sayfa = landscape(A4) if genis else A4
    belge = SimpleDocTemplate(tampon, pagesize=sayfa, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                              topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                              title=rapor["baslik"], author=okul_adi or "Faaliyet Takip")
    stil = ParagraphStyle("n", fontName="Okul", fontSize=9.5, leading=12)
    baslik = ParagraphStyle("b", fontName="Okul-Kalin", fontSize=14, leading=18, alignment=1)
    alt = ParagraphStyle("a", fontName="Okul-Kalin", fontSize=12, leading=16, alignment=1)
    hucre = ParagraphStyle("h", fontName="Okul", fontSize=8.5, leading=10.5)

    ogeler = []
    if okul_adi:
        ogeler.append(Paragraph(_kacis(okul_adi), baslik))
    ogeler.append(Paragraph(_kacis(rapor["baslik"]), alt))
    ogeler.append(Spacer(1, 6))
    for b in rapor["bilgi"]:
        ogeler.append(Paragraph(_kacis(b), stil))
    ogeler.append(Paragraph("Rapor tarihi: " + datetime.now().strftime("%d.%m.%Y %H:%M"), stil))
    ogeler.append(Spacer(1, 10))

    hizalama = {"para": 2, "yuzde": 2, "sayi": 1, "tarih": 1}
    stiller = [ParagraphStyle(f"s{i}", parent=hucre, alignment=hizalama.get(t, 0))
               for i, (_, t) in enumerate(rapor["sutunlar"])]
    veri = [[Paragraph(f"<b>{_kacis(ad)}</b>", hucre) for ad, _ in rapor["sutunlar"]]]
    for s in rapor["satirlar"]:
        veri.append([Paragraph(_kacis(hucre_metni(v, t)), st)
                     for v, (_, t), st in zip(s, rapor["sutunlar"], stiller)])
    if rapor["toplam"] is not None:
        veri.append([Paragraph(f"<b>{_kacis(hucre_metni(v, t))}</b>", st)
                     for v, (_, t), st in zip(rapor["toplam"], rapor["sutunlar"], stiller)])
    if len(veri) == 1:
        veri.append([Paragraph("Kayıt yok", hucre)] + [""] * (len(rapor["sutunlar"]) - 1))

    tablo = Table(veri, repeatRows=1, colWidths=_genislikler(rapor, belge.width))
    komutlar = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F2")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
    ]
    if rapor["toplam"] is not None and rapor["satirlar"]:
        komutlar.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DCE6F2")))
    tablo.setStyle(TableStyle(komutlar))
    ogeler.append(tablo)

    ogeler.append(Spacer(1, 36))
    imza = Table([
        ["Hazırlayan", "Onaylayan"],
        [hazirlayan, onaylayan],
        ["İmza:", "İmza:"],
    ], colWidths=[belge.width / 2] * 2)
    imza.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Okul"),
        ("FONTNAME", (0, 0), (-1, 0), "Okul-Kalin"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 2), (-1, 2), 18),
    ]))
    ogeler.append(imza)

    def sayfa_no(canvas, doc):
        canvas.setFont("Okul", 8)
        canvas.drawRightString(sayfa[0] - 1.5 * cm, 0.9 * cm, f"Sayfa {doc.page}")

    belge.build(ogeler, onFirstPage=sayfa_no, onLaterPages=sayfa_no)
    return tampon.getvalue()


def _genislikler(rapor, toplam_genislik):
    agirlik = {"metin": 2.0, "para": 1.3, "yuzde": 0.9, "tarih": 1.1, "sayi": 0.9}
    w = []
    for ad, tip in rapor["sutunlar"]:
        a = agirlik[tip]
        if ad in ("Ad Soyad", "Gösterge", "Değer", "Açıklama"):
            a = 3.0
        elif ad in ("No", "Sınıf"):
            a = 0.9
        w.append(a)
    oran = toplam_genislik / sum(w)
    return [x * oran for x in w]


def _kacis(metin):
    return str(metin).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
