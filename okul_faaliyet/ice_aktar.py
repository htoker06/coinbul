"""Sınıf ve öğrenci listesini Excel'den veya yapıştırılan metinden içe aktarır."""
import io
import re

from openpyxl import Workbook, load_workbook


def _norm(metin):
    s = str(metin or "").strip().lower()
    for a, b in (("ı", "i"), ("İ", "i"), ("ş", "s"), ("ğ", "g"), ("ü", "u"), ("ö", "o"),
                 ("ç", "c"), ("i̇", "i")):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", s)


def _sutun_bul(basliklar):
    """Başlık satırından sütun indekslerini bulur."""
    yer = {}
    for i, b in enumerate(basliklar):
        n = _norm(b)
        if not n:
            continue
        if n in ("sinif", "sinifi", "sube", "sinifsube", "sinifadi") and "sinif" not in yer:
            yer["sinif"] = i
        elif n in ("no", "okulno", "ogrencino", "numara", "okulnumarasi", "ogrencinumarasi"):
            yer.setdefault("no", i)
        elif n in ("adsoyad", "adisoyadi", "ogrenciadsoyad", "ogrenciadisoyadi", "adisoyad"):
            yer.setdefault("adsoyad", i)
        elif n in ("ad", "adi", "ogrenciadi"):
            yer.setdefault("ad", i)
        elif n in ("soyad", "soyadi"):
            yer.setdefault("soyad", i)
        elif "veli" in n or "tel" in n:
            yer.setdefault("tel", i)
    return yer


def excel_oku(icerik):
    """Excel dosyasından [(sinif, no, ad_soyad, tel), ...] döndürür.

    Sınıf sütunu yoksa sayfa adı sınıf adı kabul edilir (her sınıf ayrı sayfa).
    """
    kayitlar = []
    for sayfa_adi, satirlar in _sayfalar(icerik):
        baslik_idx, yer = None, {}
        for i, satir in enumerate(satirlar[:15]):
            y = _sutun_bul(satir)
            if ("adsoyad" in y or "ad" in y):
                baslik_idx, yer = i, y
                break
        if baslik_idx is None:
            continue
        for satir in satirlar[baslik_idx + 1:]:
            def al(k):
                return str(satir[yer[k]]).strip() if k in yer and yer[k] < len(satir) \
                    and satir[yer[k]] is not None else ""
            if "adsoyad" in yer:
                ad = al("adsoyad")
            else:
                ad = f"{al('ad')} {al('soyad')}".strip()
            if not ad:
                continue
            sinif = al("sinif") if "sinif" in yer else sayfa_adi.strip()
            no = al("no")
            if no.endswith(".0"):
                no = no[:-2]
            kayitlar.append((sinif_adi_duzelt(sinif), no, ad, al("tel")))
    return kayitlar


def _sayfalar(icerik):
    """[(sayfa adı, satırlar)]; hem .xlsx hem e-Okul'un verdiği eski .xls biçimini okur."""
    if icerik[:4] == b"\xd0\xcf\x11\xe0":  # eski Excel 97-2003 (.xls)
        import xlrd
        wb = xlrd.open_workbook(file_contents=icerik)
        return [(sh.name, [sh.row_values(r) for r in range(sh.nrows)]) for sh in wb.sheets()]
    wb = load_workbook(io.BytesIO(icerik), read_only=True, data_only=True)
    return [(ws.title, list(ws.iter_rows(values_only=True))) for ws in wb.worksheets]


def metin_oku(metin):
    """Her satır: 'no ad soyad' veya sekme/; ile ayrılmış 'no;ad soyad;veli tel'."""
    kayitlar = []
    for satir in str(metin).splitlines():
        satir = satir.strip()
        if not satir:
            continue
        if "\t" in satir or ";" in satir:
            parcalar = [p.strip() for p in re.split(r"[\t;]", satir)]
        else:
            m = re.match(r"^(\d+)[\s.)-]+(.+)$", satir)
            parcalar = [m.group(1), m.group(2)] if m else ["", satir]
        if len(parcalar) == 1:
            parcalar = ["", parcalar[0]]
        if not parcalar[0].isdigit() and parcalar[1].isdigit():
            parcalar[0], parcalar[1] = parcalar[1], parcalar[0]
        no, ad = parcalar[0], parcalar[1]
        tel = parcalar[2] if len(parcalar) > 2 else ""
        if ad:
            kayitlar.append((no, re.sub(r"\s+", " ", ad), tel))
    return kayitlar


def sinif_adi_duzelt(ad):
    """'9a', '9 / A', '9-a' -> '9-A'; tanınmayan biçim olduğu gibi kalır."""
    s = str(ad).strip()
    m = re.match(r"^(\d{1,2})\s*[./\- ]?\s*([A-Za-zÇĞİÖŞÜçğıöşü])\s*(şubesi|şube)?$", s,
                 re.IGNORECASE)
    if m:
        harf = m.group(2)
        harf = "I" if harf == "ı" else ("İ" if harf == "i" else harf.upper())
        return f"{m.group(1)}-{harf}"
    return s


def kademe_bul(sinif_adi):
    m = re.match(r"^(\d{1,2})", sinif_adi)
    return f"{m.group(1)}. sınıf" if m else ""


def sinif_getir_olustur(con, ad):
    ad = sinif_adi_duzelt(ad)
    r = con.execute("SELECT id FROM siniflar WHERE ad = ?", (ad,)).fetchone()
    if r:
        return r["id"], False
    cur = con.execute("INSERT INTO siniflar (ad, kademe) VALUES (?, ?)", (ad, kademe_bul(ad)))
    return cur.lastrowid, True


def ogrenci_kaydet(con, sinif_id, no, ad, tel=""):
    """Okul numarası varsa aynı numaralı öğrenciyi günceller (sınıf değişikliği dahil).

    Dönüş: 'yeni' veya 'guncel'.
    """
    mevcut = None
    if no:
        mevcut = con.execute("SELECT id FROM ogrenciler WHERE okul_no = ?", (no,)).fetchone()
    else:
        mevcut = con.execute(
            "SELECT id FROM ogrenciler WHERE sinif_id = ? AND ad_soyad = ? AND okul_no = ''",
            (sinif_id, ad)).fetchone()
    if mevcut:
        con.execute(
            "UPDATE ogrenciler SET sinif_id = ?, ad_soyad = ?, aktif = 1, "
            "veli_tel = CASE WHEN ? != '' THEN ? ELSE veli_tel END WHERE id = ?",
            (sinif_id, ad, tel, tel, mevcut["id"]))
        return "guncel"
    con.execute("INSERT INTO ogrenciler (sinif_id, okul_no, ad_soyad, veli_tel) VALUES (?, ?, ?, ?)",
                (sinif_id, no, ad, tel))
    return "yeni"


def toplu_kaydet(con, kayitlar):
    """[(sinif, no, ad, tel)] -> özet sözlük."""
    ozet = {"sinif": 0, "yeni": 0, "guncel": 0}
    for sinif, no, ad, tel in kayitlar:
        if not sinif:
            continue
        sid, yeni_sinif = sinif_getir_olustur(con, sinif)
        ozet["sinif"] += int(yeni_sinif)
        ozet[ogrenci_kaydet(con, sid, no, ad, tel)] += 1
    return ozet


def ornek_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Öğrenciler"
    ws.append(["Sınıf", "Okul No", "Ad Soyad", "Veli Telefonu"])
    ws.append(["9-A", "101", "Ayşe Yılmaz", "0555 111 22 33"])
    ws.append(["9-A", "102", "Mehmet Kaya", ""])
    ws.append(["10-B", "215", "Zeynep Demir", ""])
    for harf, gen in zip("ABCD", (10, 12, 28, 18)):
        ws.column_dimensions[harf].width = gen
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
