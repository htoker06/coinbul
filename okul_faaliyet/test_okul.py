import io
import os
import sys

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import db  # noqa: E402
import ice_aktar  # noqa: E402
from app import uygulama_olustur  # noqa: E402


def test_para_oku_yaz():
    assert db.para_oku("1.250,50") == 125050
    assert db.para_oku("1250,5") == 125050
    assert db.para_oku("1250.50") == 125050
    assert db.para_oku("350") == 35000
    assert db.para_oku("1.500") == 150000
    assert db.para_oku("350 TL") == 35000
    with pytest.raises(ValueError):
        db.para_oku("abc")
    assert db.para_yaz(125050) == "1.250,50 TL"
    assert db.para_yaz(0) == "0,00 TL"


def test_tarih():
    assert db.tarih_oku("15.10.2026") == "2026-10-15"
    assert db.tarih_oku("2026-10-15") == "2026-10-15"
    assert db.tarih_yaz("2026-10-15") == "15.10.2026"


def test_sinif_adi():
    assert ice_aktar.sinif_adi_duzelt("9a") == "9-A"
    assert ice_aktar.sinif_adi_duzelt("10 / b") == "10-B"
    assert ice_aktar.sinif_adi_duzelt("9-i") == "9-İ"
    assert ice_aktar.sinif_adi_duzelt("Anasınıfı") == "Anasınıfı"
    assert sorted(["10-A", "9-B", "9-A"], key=db.sinif_sira_anahtari) == ["9-A", "9-B", "10-A"]


def test_metin_oku():
    k = ice_aktar.metin_oku("101 Ayşe Yılmaz\n102\tMehmet Kaya\t0555\n\nZeynep Demir")
    assert k == [("101", "Ayşe Yılmaz", ""), ("102", "Mehmet Kaya", "0555"),
                 ("", "Zeynep Demir", "")]


@pytest.fixture
def istemci(tmp_path):
    app = uygulama_olustur(str(tmp_path / "okul.db"))
    app.config["TESTING"] = True
    c = app.test_client()
    c.app = app
    return c


def _csrf(c):
    with c.session_transaction() as s:
        s.setdefault("csrf", "test")
        return s["csrf"]


def post(c, url, **veri):
    veri["csrf"] = _csrf(c)
    return c.post(url, data=veri, follow_redirects=True)


def _excel(satirlar):
    wb = Workbook()
    ws = wb.active
    for s in satirlar:
        ws.append(s)
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    return b


def kurulmus(c):
    assert c.get("/").status_code == 302
    r = post(c, "/kurulum", okul_adi="Deneme Lisesi", ad="Ali Veli", kullanici_adi="mudur",
             sifre="123456", sifre2="123456")
    assert "Kurulum tamamlandı" in r.get_data(as_text=True)
    dosya = _excel([["Sınıf", "Okul No", "Ad Soyad", "Veli Telefonu"],
                    ["9a", 101, "Ayşe Yılmaz", "0555"],
                    ["9-A", 102, "Mehmet Kaya", ""],
                    ["9-A", 103, "Can Ak", ""],
                    ["10-B", 201, "Zeynep Demir", ""]])
    r = c.post("/ice-aktar", data={"csrf": _csrf(c), "dosya": (dosya, "liste.xlsx")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "2 yeni sınıf, 4 yeni öğrenci" in r.get_data(as_text=True)


def _idler(c):
    con = db.baglan(c.app.config["VERITABANI"])
    sinif = {r["ad"]: r["id"] for r in con.execute("SELECT * FROM siniflar")}
    ogr = {r["ad_soyad"]: r["id"] for r in con.execute("SELECT * FROM ogrenciler")}
    con.close()
    return sinif, ogr


def test_tam_akis(istemci):
    c = istemci
    kurulmus(c)
    sinif, ogr = _idler(c)
    r = post(c, "/faaliyet/yeni", ad="Çanakkale gezisi", ucret="350", son_tarih="2026-11-01",
             sinif=[sinif["9-A"], sinif["10-B"]])
    assert "oluşturuldu" in r.get_data(as_text=True)
    fid = 1

    # Kısmi ödeme, sonra tamamlama
    r = post(c, f"/faaliyet/{fid}/odeme", ogrenci_id=ogr["Ayşe Yılmaz"], tutar="150",
             tarih="2026-10-10")
    assert "Kalan: 200,00 TL" in r.get_data(as_text=True)
    r = post(c, f"/faaliyet/{fid}/odeme", ogrenci_id=ogr["Ayşe Yılmaz"], tutar="200,00",
             tarih="2026-10-12")
    assert "TAMAMLANDI" in r.get_data(as_text=True)
    post(c, f"/faaliyet/{fid}/odeme", ogrenci_id=ogr["Mehmet Kaya"], tutar="100", tarih="2026-10-12")

    # Fazla ödeme onaysız reddedilir
    r = post(c, f"/faaliyet/{fid}/odeme", ogrenci_id=ogr["Zeynep Demir"], tutar="500",
             tarih="2026-10-12")
    assert "fazla" in r.get_data(as_text=True)

    # Can muaf
    post(c, f"/faaliyet/{fid}/ogrenci/{ogr['Can Ak']}/ayar", tip="muaf", aciklama="şehit yakını")

    con = db.baglan(c.app.config["VERITABANI"])
    durum = {o["ad_soyad"]: o for o in db.ogrenci_durumlari(con, fid)}
    assert durum["Ayşe Yılmaz"]["durum"] == db.DURUM_TAMAM
    assert durum["Mehmet Kaya"]["durum"] == db.DURUM_KISMI
    assert durum["Mehmet Kaya"]["kalan"] == 25000
    assert durum["Zeynep Demir"]["durum"] == db.DURUM_YOK
    assert durum["Can Ak"]["durum"] == db.DURUM_MUAF
    oz = {s["sinif"]: s for s in db.sinif_ozetleri(con, fid)}
    assert oz["9-A"]["hedef"] == 70000 and oz["9-A"]["toplanan"] == 45000
    assert oz["10-B"]["toplanan"] == 0
    # Ödeme iptali
    pid = con.execute("SELECT id FROM odemeler WHERE ogrenci_id = ?",
                      (ogr["Mehmet Kaya"],)).fetchone()[0]
    con.close()
    post(c, f"/odeme/{pid}/iptal", neden="yanlış öğrenci")
    con = db.baglan(c.app.config["VERITABANI"])
    m = db.ogrenci_durumlari(con, fid, ogrenci_id=ogr["Mehmet Kaya"])[0]
    assert m["durum"] == db.DURUM_YOK
    assert con.execute("SELECT COUNT(*) FROM islem_gecmisi WHERE islem = 'Ödeme iptal edildi'"
                       ).fetchone()[0] == 1
    con.close()

    # Sayfalar açılıyor
    for url in ["/", f"/faaliyet/{fid}", f"/faaliyet/{fid}/sinif/{sinif['9-A']}",
                f"/faaliyet/{fid}/ogrenci/{ogr['Ayşe Yılmaz']}", "/siniflar",
                f"/sinif/{sinif['9-A']}", "/kullanicilar", "/gecmis", "/ayarlar", "/ice-aktar"]:
        assert c.get(url).status_code == 200, url

    # Raporlar: ekran, Excel, PDF
    for tur in ["sinif_listesi", "odemeyenler", "eksik", "sinif_ozet", "okul_ozet", "tahsilat"]:
        r = c.get(f"/raporlar?faaliyet={fid}&tur={tur}")
        assert r.status_code == 200, tur
        x = c.get(f"/raporlar?faaliyet={fid}&tur={tur}&bicim=xlsx")
        assert x.status_code == 200 and x.data[:2] == b"PK", tur
        load_workbook(io.BytesIO(x.data))
        p = c.get(f"/raporlar?faaliyet={fid}&tur={tur}&bicim=pdf")
        assert p.status_code == 200 and p.data[:4] == b"%PDF", tur

    html = c.get(f"/raporlar?faaliyet={fid}&tur=odemeyenler").get_data(as_text=True)
    assert "Zeynep Demir" in html and "Mehmet Kaya" in html and "Ayşe Yılmaz" not in html

    ws = load_workbook(io.BytesIO(
        c.get(f"/raporlar?faaliyet={fid}&tur=sinif_ozet&bicim=xlsx").data)).active
    degerler = [[h.value for h in s] for s in ws.iter_rows()]
    satir_9a = next(s for s in degerler if s[0] == "9-A")
    assert satir_9a[2] == 700 and satir_9a[3] == 350


def test_ogretmen_yetkisi(istemci):
    c = istemci
    kurulmus(c)
    sinif, ogr = _idler(c)
    post(c, "/faaliyet/yeni", ad="Tiyatro", ucret="100", sinif=[sinif["9-A"], sinif["10-B"]])
    post(c, "/kullanicilar", ad="Öğretmen Bey", kullanici_adi="ogr", sifre="654321",
         rol="ogretmen", sinif_id=sinif["9-A"])
    post(c, "/cikis")
    r = post(c, "/giris", kullanici_adi="ogr", sifre="654321")
    assert r.status_code == 200
    # Kendi sınıfına ödeme girer, başka sınıfa giremez
    r = post(c, "/faaliyet/1/odeme", ogrenci_id=ogr["Ayşe Yılmaz"], tutar="100", tarih="2026-10-01")
    assert "TAMAMLANDI" in r.get_data(as_text=True)
    r = c.post("/faaliyet/1/odeme", data={"csrf": _csrf(c), "ogrenci_id": ogr["Zeynep Demir"],
                                          "tutar": "100", "tarih": "2026-10-01"})
    assert r.status_code == 403
    assert c.get(f"/faaliyet/1/sinif/{sinif['10-B']}").status_code == 403
    assert c.get("/siniflar").status_code == 403
    html = c.get("/raporlar?faaliyet=1&tur=sinif_listesi").get_data(as_text=True)
    assert "Ayşe Yılmaz" in html and "Zeynep Demir" not in html


def test_csrf_zorunlu(istemci):
    kurulmus(istemci)
    r = istemci.post("/faaliyet/yeni", data={"ad": "x", "ucret": "1"})
    assert r.status_code == 400


def test_kapali_faaliyete_odeme_yok(istemci):
    c = istemci
    kurulmus(c)
    sinif, ogr = _idler(c)
    post(c, "/faaliyet/yeni", ad="Gezi", ucret="100", sinif=[sinif["9-A"]])
    post(c, "/faaliyet/1/duzenle", ad="Gezi", ucret="100", sinif=[sinif["9-A"]], kapali="1")
    r = post(c, "/faaliyet/1/odeme", ogrenci_id=ogr["Ayşe Yılmaz"], tutar="50", tarih="2026-10-01")
    assert "kapatılmış" in r.get_data(as_text=True)


def test_ayni_numara_tekrar_yuklenince_guncellenir(istemci):
    c = istemci
    kurulmus(c)
    dosya = _excel([["Okul No", "Adı", "Soyadı"], [101, "Ayşe", "Yılmaz Kara"]])
    wb = load_workbook(dosya)
    wb.active.title = "10-B"
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    r = c.post("/ice-aktar", data={"csrf": _csrf(c), "dosya": (b, "l.xlsx")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "0 yeni öğrenci, 1 güncellenen" in r.get_data(as_text=True)
    con = db.baglan(c.app.config["VERITABANI"])
    o = con.execute("SELECT o.ad_soyad, s.ad FROM ogrenciler o JOIN siniflar s ON s.id = o.sinif_id "
                    "WHERE okul_no = '101'").fetchone()
    assert (o[0], o[1]) == ("Ayşe Yılmaz Kara", "10-B")
